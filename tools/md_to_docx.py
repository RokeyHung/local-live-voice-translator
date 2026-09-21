# /// script
# requires-python = ">=3.11"
# dependencies = ["python-docx>=1.1.2"]
# ///
"""Chuyển tài liệu Markdown của ``docs/`` sang Word (.docx) để in và gửi thầy.

    uv run --no-project tools/md_to_docx.py docs/gvhd/bao-cao-bo-danh-gia.md \
        --out docs/word/bao-cao-bo-danh-gia.docx --toc

**Vì sao là script riêng có khối PEP 723** — giống ``eval_comet.py``: ``python-docx`` chỉ
dùng cho việc xuất tài liệu, không phải phụ thuộc của dịch vụ AI. Thêm nó vào
``apps/ai-service`` sẽ kéo một thư viện vô can vào môi trường chạy model, còn ``uv sync``
thì đồng bộ về **đúng** những gì lệnh nêu ra nên nó lại là một chỗ nữa để lệch môi
trường. Khối metadata ở đầu file khiến ``uv run --no-project`` tự dựng môi trường riêng.

**Vì sao không dùng pandoc.** Pandoc làm được việc này, nhưng định dạng thì phải đi qua
một ``reference.docx`` dựng bằng tay trong Word — tức là một file nhị phân không đọc được
bằng ``git diff``. Ở đây quy cách trình bày (font, giãn dòng, lề, cỡ chữ bảng) nằm trong
hằng số ở đầu file này, sửa được và xem lịch sử được.

Phạm vi Markdown đã hiện thực là **đúng những gì ``docs/`` đang dùng**, không hơn: tiêu
đề, đoạn văn, bảng GFM (kèm căn lề theo dòng ``---:``), danh sách có lồng và ô
``- [x]``, khối trích dẫn, khối mã, đường kẻ ngang, và công thức ``$...$``/``$$...$$``.
Chi tiết về công thức: xem ``latex_to_text``.
"""

from __future__ import annotations

import argparse
import re
import sys
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.text.paragraph import Paragraph

# --- Quy cách trình bày -------------------------------------------------------------
# Mặc định theo lối trình bày đồ án: Times New Roman 13pt, giãn dòng 1,5, lề 3-2-2-2 cm.
BODY_FONT = "Times New Roman"
BODY_SIZE = Pt(13)
MONO_FONT = "Consolas"
LINE_SPACING = 1.5
MARGIN_TOP, MARGIN_BOTTOM, MARGIN_LEFT, MARGIN_RIGHT = Cm(2), Cm(2), Cm(3), Cm(2)
HEADING_SIZES = {1: Pt(16), 2: Pt(14), 3: Pt(13), 4: Pt(13), 5: Pt(13), 6: Pt(13)}
# Bảng của docs/ có cái tới 11 cột (bảng độ trễ sáu chiều). Cỡ chữ giảm dần theo số cột
# để bảng không bị Word ép xuống dòng ở từng ô.
TABLE_SIZES = ((6, Pt(11)), (8, Pt(10)), (10, Pt(9)), (99, Pt(8)))
CODE_SIZE = Pt(9)
SHADE_CODE = "F2F2F2"
SHADE_TABLE_HEADER = "E8E8E8"
QUOTE_BAR = "9E9E9E"

# Chế độ --uit: quy định "Hình thức trình bày khóa luận tốt nghiệp" (Phụ lục 2) của Phòng
# Đào tạo Đại học UIT, bản 03/2024. Chỉ đổi đúng những gì quy định nêu; phần còn lại giữ
# như chế độ mặc định.
UIT_MARGINS = (Cm(3), Cm(3.5), Cm(3.5), Cm(2))  # trên, dưới, trái, phải
UIT_HEADING_SIZES = {1: Pt(14), 2: Pt(13), 3: Pt(13), 4: Pt(13), 5: Pt(13), 6: Pt(13)}
UIT = False  # bật bằng --uit
# Bề rộng vùng chữ = khổ A4 21 cm trừ lề trái 3,5 và lề phải 2.
UIT_TEXT_WIDTH = Cm(15.5)
CAPTION_STYLES = {"Hình": "Chú thích hình", "Bảng": "Chú thích bảng"}
RE_CAPTION = re.compile(r"^(Hình|Bảng) \d+\.\d+:")
RE_IMAGE = re.compile(r"^\s*!\[(?P<alt>[^\]]*)\]\((?P<src>[^)]+)\)\s*$")
# Cỡ chữ trang bìa theo mẫu bìa của trường: tên đề tài 18–30, tên tiếng Anh 15–25,
# trường/khoa/"KHÓA LUẬN TỐT NGHIỆP" 16, năm 13, còn lại 14. Đoán theo nội dung dòng
# vì Markdown không có chỗ ghi cỡ chữ.
COVER_SIZES = (
    (re.compile(r"^(TRƯỜNG|KHOA|KHÓA LUẬN)"), Pt(16)),
    (re.compile(r"^TP\. HỒ CHÍ MINH"), Pt(13)),
    (re.compile(r"^[A-Z][a-z].*[A-Za-z]$"), Pt(16)),  # tên tiếng Anh
    (re.compile(r"^[A-ZÀ-Ỹ ]{40,}$"), Pt(20)),  # tên đề tài: dòng chữ hoa dài
)

# --- Phân tích Markdown -------------------------------------------------------------

RE_HEADING = re.compile(r"^(#{1,6})\s+(.*)$")
RE_FENCE = re.compile(r"^\s*(`{3,}|~{3,})\s*([A-Za-z0-9_+-]*)\s*$")
RE_HR = re.compile(r"^\s*(-{3,}|\*{3,}|_{3,})\s*$")
RE_TABLE_SEP = re.compile(r"^\s*\|?[\s:|-]+\|[\s:|-]*$")
RE_LIST = re.compile(r"^(\s*)(?:([-*+])|(\d{1,3})[.)])\s+(.*)$")
RE_QUOTE = re.compile(r"^\s*>\s?(.*)$")
RE_TASK = re.compile(r"^\[([ xX])\]\s+(.*)$")
RE_DISPLAY_MATH = re.compile(r"^\s*\$\$(.*)\$\$\s*$")
RE_META = re.compile(r"^\*\*[^*]+:\*\*")


@dataclass
class Block:
    """Một khối Markdown đã tách xong, chờ đổ vào Word."""

    kind: str  # heading | para | table | list | quote | code | math | hr | image | directive
    text: str = ""
    level: int = 0
    lang: str = ""
    lines: list[str] | None = None
    rows: list[list[str]] | None = None
    aligns: list[str] | None = None
    items: list[tuple[int, str, str]] | None = None  # (mức lồng, dấu đầu dòng, nội dung)


def split_row(line: str) -> list[str]:
    """Tách một dòng bảng GFM, tôn trọng ``\\|`` trong ô."""
    cells = re.split(r"(?<!\\)\|", line.strip())
    if cells and not cells[0].strip():
        cells.pop(0)
    if cells and not cells[-1].strip():
        cells.pop()
    return [c.strip().replace(r"\|", "|") for c in cells]


def parse_aligns(sep: str) -> list[str]:
    aligns = []
    for cell in split_row(sep):
        left, right = cell.startswith(":"), cell.endswith(":")
        aligns.append("center" if left and right else "right" if right else "left")
    return aligns


def parse(md: str) -> list[Block]:
    """Markdown → danh sách khối. Một lượt duyệt, không cây cú pháp — đủ cho ``docs/``."""
    lines = md.replace("\r\n", "\n").split("\n")
    blocks: list[Block] = []
    buf: list[str] = []  # đoạn văn đang gom (Markdown nối dòng mềm thành một đoạn)

    def flush() -> None:
        if not buf:
            return
        # Khối thông tin đầu tài liệu ("**Đề tài:** … / **Sinh viên:** …") là nhiều dòng
        # liền nhau, mỗi dòng một nhãn. Markdown nối chúng thành một đoạn, nhưng trên
        # trang giấy thì phải xuống dòng, nên tách lại khi thấy từ hai nhãn trở lên.
        labels = [n for n, s in enumerate(buf) if RE_META.match(s.strip())]
        if len(labels) >= 2:
            for start, end in zip(labels, labels[1:] + [len(buf)]):
                blocks.append(Block("para", text=" ".join(s.strip() for s in buf[start:end])))
        else:
            blocks.append(Block("para", text=" ".join(s.strip() for s in buf).strip()))
        buf.clear()

    i = 0
    while i < len(lines):
        line = lines[i]

        # Chú thích HTML không phải nội dung. Trước đây nó bị in ra thành chữ; giờ nó
        # thành một khối "directive" — chế độ --uit đọc vài chỉ thị trong đó (ngắt
        # trang, mục lục, bắt đầu đánh số trang), chế độ mặc định bỏ qua.
        if line.lstrip().startswith("<!--"):
            flush()
            body = [line]
            while "-->" not in body[-1] and i + 1 < len(lines):
                i += 1
                body.append(lines[i])
            text = " ".join(s.strip() for s in body)
            blocks.append(Block("directive", text=text[4 : text.rfind("-->")].strip()))
            i += 1
            continue

        if m := RE_IMAGE.match(line):
            flush()
            blocks.append(Block("image", text=m.group("alt"), lines=[m.group("src")]))
            i += 1
            continue

        if fence := RE_FENCE.match(line):
            flush()
            closing, lang = fence.group(1)[0], fence.group(2)
            body: list[str] = []
            i += 1
            while i < len(lines) and not re.match(rf"^\s*{closing}{{3,}}\s*$", lines[i]):
                body.append(lines[i])
                i += 1
            blocks.append(Block("code", lang=lang, lines=body))
            i += 1
            continue

        if not line.strip():
            flush()
            i += 1
            continue

        if m := RE_DISPLAY_MATH.match(line):
            flush()
            blocks.append(Block("math", text=m.group(1).strip()))
            i += 1
            continue

        if m := RE_HEADING.match(line):
            flush()
            blocks.append(Block("heading", text=m.group(2).strip(), level=len(m.group(1))))
            i += 1
            continue

        # Đường kẻ ngang phải xét sau danh sách: `---` và `- x` dễ lẫn nhau.
        if RE_HR.match(line) and not RE_LIST.match(line):
            flush()
            blocks.append(Block("hr"))
            i += 1
            continue

        # Bảng: dòng hiện tại có `|` và dòng sau là dòng phân cách.
        if "|" in line and i + 1 < len(lines) and RE_TABLE_SEP.match(lines[i + 1]):
            flush()
            header = split_row(line)
            aligns = parse_aligns(lines[i + 1])
            rows = [header]
            i += 2
            while i < len(lines) and "|" in lines[i] and lines[i].strip():
                rows.append(split_row(lines[i]))
                i += 1
            width = max(len(r) for r in rows)
            rows = [r + [""] * (width - len(r)) for r in rows]
            aligns = (aligns + ["left"] * width)[:width]
            blocks.append(Block("table", rows=rows, aligns=aligns))
            continue

        if RE_QUOTE.match(line):
            flush()
            quoted: list[str] = []
            while i < len(lines) and (m := RE_QUOTE.match(lines[i])):
                quoted.append(m.group(1))
                i += 1
            # Trong khối trích dẫn, dòng trống ngăn đoạn; còn lại là nối dòng mềm.
            paras, cur = [], []
            for q in quoted:
                if q.strip():
                    cur.append(q.strip())
                elif cur:
                    paras.append(" ".join(cur))
                    cur = []
            if cur:
                paras.append(" ".join(cur))
            blocks.append(Block("quote", lines=paras))
            continue

        if RE_LIST.match(line):
            flush()
            items: list[tuple[int, str, str]] = []
            while i < len(lines):
                if m := RE_LIST.match(lines[i]):
                    indent = len(m.group(1).expandtabs(4))
                    marker = f"{m.group(3)}." if m.group(3) else "•"
                    items.append((indent // 2, marker, m.group(4).strip()))
                    i += 1
                elif lines[i].strip() and lines[i][:1] in " \t":
                    # Dòng tiếp nối của ô danh sách phía trên.
                    lvl, marker, txt = items[-1]
                    items[-1] = (lvl, marker, f"{txt} {lines[i].strip()}")
                    i += 1
                else:
                    break
            # Mức lồng trong docs/ thụt 2 hoặc 3 dấu cách tuỳ dấu đầu dòng; chuẩn hoá về
            # 0,1,2… theo thứ tự các mức thụt đã gặp, nếu không "  -" và "   -" sẽ ra hai
            # mức khác nhau dù cùng là con của một ô.
            seen = sorted({lvl for lvl, _, _ in items})
            depth = {lvl: n for n, lvl in enumerate(seen)}
            blocks.append(Block("list", items=[(depth[lv], mk, tx) for lv, mk, tx in items]))
            continue

        buf.append(line)
        i += 1

    flush()
    return blocks


# --- Công thức ----------------------------------------------------------------------

SUB = str.maketrans("0123456789+-=()naeimnt", "₀₁₂₃₄₅₆₇₈₉₊₋₌₍₎ₙₐₑᵢₘₙₜ")
SUP = str.maketrans("0123456789+-=()n", "⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻⁼⁽⁾ⁿ")
SYMBOLS = {
    r"\cdot": "·",
    r"\times": "×",
    r"\div": "÷",
    r"\le": "≤",
    r"\leq": "≤",
    r"\ge": "≥",
    r"\geq": "≥",
    r"\approx": "≈",
    r"\neq": "≠",
    r"\pm": "±",
    r"\infty": "∞",
    r"\sum": "Σ",
    r"\prod": "Π",
    r"\alpha": "α",
    r"\beta": "β",
    r"\lambda": "λ",
    r"\mu": "μ",
    r"\sigma": "σ",
    r"\to": "→",
    r"\rightarrow": "→",
    r"\quad": "   ",
    r"\qquad": "     ",
    r"\,": " ",
    r"\;": " ",
    r"\!": "",
    r"\left": "",
    r"\right": "",
    r"\%": "%",
    r"\\": "; ",
}
# Lệnh LaTeX phải khớp cả tên mới thay, không thì `\le` ăn mất hai chữ đầu của `\left`.
RE_TEX_CMD = re.compile(r"\\[A-Za-z]+|\\[\\,;!%]")


def wrap_operand(part: str) -> str:
    """Thêm ngoặc cho tử/mẫu khi nó là một tổng — ``(S + D + I) / N``, không phải
    ``S + D + I / N`` (đọc ra thành ``S + D + (I/N)``). Cụm chữ thường thì không thêm."""
    # Phải là dấu cộng/trừ có khoảng trắng hai bên, không thì "wall-clock" cũng bị ngoặc.
    text = part.strip()
    return f"({text})" if re.search(r"\s[+\-−±]\s", text) else text


def latex_to_text(src: str) -> str:
    """LaTeX → một dòng chữ đọc được.

    Không sinh OMML (công thức "thật" của Word): cả tài liệu chỉ có bốn công thức, mà
    dựng OMML bằng tay thì dài hơn toàn bộ script này. Đổi lấy chữ Unicode một dòng —
    ``\\frac{a}{b}`` thành ``a / b``, chỉ số dưới/trên thành ký tự Unicode — nên công
    thức vẫn đọc đúng trên Word, chỉ không sửa được bằng trình soạn công thức.
    """
    s = src.strip()
    s = re.sub(r"\\(?:text|mathrm|mathbf|operatorname)\s*\{([^{}]*)\}", r"\1", s)
    s = re.sub(r"\\begin\{cases\}(.*?)\\end\{cases\}", r"\1", s, flags=re.S)
    # \frac lồng nhau: chạy lại tới khi hết, mỗi lượt xử lý cái trong cùng.
    for _ in range(4):
        new = re.sub(
            r"\\[dt]?frac\s*\{([^{}]*)\}\s*\{([^{}]*)\}",
            lambda m: f"{wrap_operand(m.group(1))} / {wrap_operand(m.group(2))}",
            s,
        )
        if new == s:
            break
        s = new
    s = re.sub(r"_\{([^{}]*)\}", lambda m: m.group(1).translate(SUB), s)
    s = re.sub(r"\^\{([^{}]*)\}", lambda m: m.group(1).translate(SUP), s)
    s = re.sub(r"_([A-Za-z0-9])", lambda m: m.group(1).translate(SUB), s)
    s = re.sub(r"\^([A-Za-z0-9])", lambda m: m.group(1).translate(SUP), s)
    # Lệnh không có trong bảng (\exp, \log) thì bỏ dấu gạch chéo, giữ lại tên.
    s = RE_TEX_CMD.sub(lambda m: SYMBOLS.get(m.group(0), m.group(0).lstrip("\\")), s)
    s = s.replace("&", " ").replace("{,}", ",").replace("{", "").replace("}", "")
    return re.sub(r"\s{2,}", lambda m: m.group(0) if len(m.group(0)) > 2 else " ", s).strip()


# --- Chữ trong dòng ------------------------------------------------------------------

RE_INLINE = re.compile(
    r"(?P<ticks>`+)(?P<code>.+?)(?P=ticks)"
    r"|!\[(?P<alt>[^\]]*)\]\((?P<src>[^)]*)\)"
    r"|\[(?P<ltext>[^\]]+)\]\((?P<lurl>[^)\s]+)(?:\s+\"[^\"]*\")?\)"
    r"|<(?P<autobr>https?://[^>\s]+)>"
    r"|(?P<auto>https?://[^\s<>()\[\]]+)"
    r"|\*\*(?P<bold>.+?)\*\*"
    r"|__(?P<bold2>.+?)__"
    r"|\$(?P<math>[^$\n]+)\$"
    r"|\*(?P<it>[^*\n]+)\*"
    r"|(?<![0-9A-Za-zÀ-ỹ])_(?P<it2>[^_\n]+)_(?![0-9A-Za-zÀ-ỹ])",
    re.S,
)


SOURCE_DIR = Path(".")  # thư mục của file .md đang xử lý, để lần theo liên kết nội bộ
DOC_TITLES: dict[Path, str] = {}
RE_DOC_REF = re.compile(r"^(\d{2})(\s+.+)?$")


def doc_title(target: str) -> str:
    """Tiêu đề (H1) của một file .md trong ``docs/``, rỗng nếu không có.

    Dùng để thay những nhãn liên kết chỉ có số thứ tự — ``xem [06](06_….md)`` — bằng tên
    thật của tài liệu. Trên bản Word người đọc không có cây file để tự suy ra "06" là gì.
    """
    path = (SOURCE_DIR / target.split("#")[0]).resolve()
    if path not in DOC_TITLES:
        title = ""
        if path.is_file() and path.suffix == ".md":
            for line in path.read_text(encoding="utf-8").splitlines():
                if line.startswith("# "):
                    title = line[2:].split(":")[0].strip()
                    break
        DOC_TITLES[path] = title if len(title) <= 60 else ""
    return DOC_TITLES[path]


@dataclass
class Frag:
    text: str
    bold: bool = False
    italic: bool = False
    mono: bool = False
    url: str = ""


def inline(text: str, bold: bool = False, italic: bool = False) -> list[Frag]:
    """Tách chữ trong dòng thành các đoạn mang định dạng, đệ quy một cấp cho lồng nhau."""
    out: list[Frag] = []
    pos = 0
    for m in RE_INLINE.finditer(text):
        if m.start() > pos:
            out.append(Frag(unescape(text[pos : m.start()]), bold, italic))
        if m.group("code") is not None:
            out.append(Frag(m.group("code"), bold, italic, mono=True))
        elif m.group("alt") is not None:  # docs/ không có ảnh; giữ chú thích cho chắc
            out.append(Frag(f"[hình: {m.group('alt')}]", bold, italic=True))
        elif m.group("ltext") is not None:
            url = m.group("lurl")
            label = m.group("ltext")
            if url.startswith(("http://", "https://")):
                out.append(Frag(label, bold, italic, url=url))
            else:
                # Liên kết nội bộ giữa các file .md: bản Word không có file kia, nên chỉ
                # giữ phần chữ — thay vì để lại một đường dẫn chết cho người đọc bản in.
                ref = RE_DOC_REF.match(label.replace("`", "").strip())
                if ref and (title := doc_title(url)):
                    label = f"{title}{ref.group(2) or ''}"
                out.extend(inline(label, bold, italic))
        elif m.group("autobr") is not None:
            # Autolink `<https://…>`: giữ URL hiện nguyên văn như trong danh sách nguồn,
            # nhưng bỏ cặp ngoặc nhọn — chúng là cú pháp Markdown, không phải nội dung.
            url = m.group("autobr")
            out.append(Frag(url, bold, italic, url=url))
        elif m.group("auto") is not None:
            # Danh sách nguồn trích dẫn viết URL trần; trên Word nên bấm được.
            url = m.group("auto").rstrip(".,;:")
            out.append(Frag(url, bold, italic, url=url))
            pos = m.start() + len(url)
            continue
        elif m.group("bold") is not None:
            out.extend(inline(m.group("bold"), True, italic))
        elif m.group("bold2") is not None:
            out.extend(inline(m.group("bold2"), True, italic))
        elif m.group("math") is not None:
            out.append(Frag(latex_to_text(m.group("math")), bold, italic=True))
        elif m.group("it") is not None:
            out.extend(inline(m.group("it"), bold, True))
        elif m.group("it2") is not None:
            out.extend(inline(m.group("it2"), bold, True))
        pos = m.end()
    if pos < len(text):
        out.append(Frag(unescape(text[pos:]), bold, italic))
    return out


def unescape(s: str) -> str:
    return re.sub(r"\\([\\`*_{}\[\]()#+\-.!|~])", r"\1", s)


# --- Dựng file Word ------------------------------------------------------------------


def shade(element, color: str) -> None:
    from docx.oxml import OxmlElement

    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), color)
    element.append(shd)


def set_font(run, name: str) -> None:
    """Đặt font cho cả ba nhóm ký tự — chỉ đặt ``run.font.name`` thì chữ có dấu tiếng
    Việt vẫn rơi về font mặc định của Word trên một số bản."""
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        from docx.oxml import OxmlElement

        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        fonts.set(qn(attr), name)


def add_hyperlink(paragraph: Paragraph, text: str, url: str) -> None:
    from docx.oxml import OxmlElement

    rel = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(color)
    rpr.append(underline)
    run.append(rpr)
    node = OxmlElement("w:t")
    node.text = text
    node.set(qn("xml:space"), "preserve")
    run.append(node)
    link.append(run)
    paragraph._p.append(link)


def add_field(paragraph: Paragraph, instruction: str, placeholder: str) -> None:
    """Chèn một field của Word (PAGE, TOC). Word tính giá trị khi mở/cập nhật file."""
    from docx.oxml import OxmlElement

    begin = OxmlElement("w:r")
    char = OxmlElement("w:fldChar")
    char.set(qn("w:fldCharType"), "begin")
    begin.append(char)
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    run_instr = OxmlElement("w:r")
    run_instr.append(instr)
    sep = OxmlElement("w:r")
    char_sep = OxmlElement("w:fldChar")
    char_sep.set(qn("w:fldCharType"), "separate")
    sep.append(char_sep)
    text_run = OxmlElement("w:r")
    node = OxmlElement("w:t")
    node.text = placeholder
    text_run.append(node)
    end = OxmlElement("w:r")
    char_end = OxmlElement("w:fldChar")
    char_end.set(qn("w:fldCharType"), "end")
    end.append(char_end)
    for el in (begin, run_instr, sep, text_run, end):
        paragraph._p.append(el)


def add_fragments(paragraph: Paragraph, frags: list[Frag], size: Pt | None = None) -> None:
    for frag in frags:
        if frag.url:
            add_hyperlink(paragraph, frag.text, frag.url)
            continue
        if not frag.text:
            continue
        run = paragraph.add_run(frag.text)
        run.bold = frag.bold
        run.italic = frag.italic
        set_font(run, MONO_FONT if frag.mono else BODY_FONT)
        run.font.size = (size or BODY_SIZE) - (Pt(1) if frag.mono else Pt(0))


def base_document() -> Document:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = BODY_SIZE
    normal.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), BODY_FONT)
    pf = normal.paragraph_format
    pf.line_spacing = LINE_SPACING
    pf.space_after = Pt(6)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for level, size in HEADING_SIZES.items():
        style = doc.styles[f"Heading {level}"]
        style.font.name = BODY_FONT
        style.font.size = size
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), BODY_FONT)
        hpf = style.paragraph_format
        hpf.space_before = Pt(12 if level <= 2 else 8)
        hpf.space_after = Pt(6)
        hpf.line_spacing = 1.3
        hpf.keep_with_next = True
        hpf.alignment = WD_ALIGN_PARAGRAPH.LEFT

    section = doc.sections[0]
    if UIT:
        section.page_width, section.page_height = Cm(21), Cm(29.7)
        top, bottom, left, right = UIT_MARGINS
    else:
        top, bottom, left, right = MARGIN_TOP, MARGIN_BOTTOM, MARGIN_LEFT, MARGIN_RIGHT
    section.top_margin, section.bottom_margin = top, bottom
    section.left_margin, section.right_margin = left, right

    if UIT:
        # Hai kiểu chú thích riêng cho hình và bảng, để Danh mục hình và Danh mục bảng là
        # hai field TOC lọc theo từng kiểu — Word tự điền số trang.
        from docx.enum.style import WD_STYLE_TYPE

        for name in CAPTION_STYLES.values():
            style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            style.base_style = normal
            style.font.size = Pt(12)
            style.font.italic = True
            spf = style.paragraph_format
            spf.alignment = WD_ALIGN_PARAGRAPH.CENTER
            spf.space_before, spf.space_after = Pt(3), Pt(9)
            spf.line_spacing = 1.2
        # Trang bìa, lời cảm ơn, mục lục, danh mục không đánh số (quy định); số trang bắt
        # đầu từ Tóm tắt, ở section thứ hai — xem start_numbering().
        request_field_update(doc)
        return doc

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(footer, " PAGE ", "1")
    for run in footer.runs:
        set_font(run, BODY_FONT)
        run.font.size = Pt(11)
    return doc


def request_field_update(doc: Document) -> None:
    """Bảo Word cập nhật mọi field (mục lục, danh mục, số trang) khi mở file.

    Không có cờ này thì mục lục hiện dòng nhắc cho tới khi người dùng tự bấm F9 — dễ nộp
    nhầm một bản chưa có mục lục.
    """
    from docx.oxml import OxmlElement

    flag = OxmlElement("w:updateFields")
    flag.set(qn("w:val"), "true")
    doc.settings.element.append(flag)


def start_numbering(doc: Document) -> None:
    """Mở section mới ở trang mới, đánh số trang từ 1 ở giữa chân trang."""
    from docx.enum.section import WD_SECTION
    from docx.oxml import OxmlElement

    section = doc.add_section(WD_SECTION.NEW_PAGE)
    section.footer.is_linked_to_previous = False
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(footer, " PAGE ", "1")
    for run in footer.runs:
        set_font(run, BODY_FONT)
        run.font.size = Pt(12)
    pg = OxmlElement("w:pgNumType")
    pg.set(qn("w:start"), "1")
    section._sectPr.append(pg)


def add_picture_block(doc: Document, src: str, alt: str) -> None:
    """Chèn ảnh canh giữa, co về bề rộng vùng chữ nếu ảnh rộng hơn.

    Ảnh chưa có (ảnh chụp màn hình phải chụp tay) thì để một ô nhắc thay vì dừng cả lượt
    xuất — bản nháp vẫn xuất được, và chỗ còn thiếu hiện ra rõ ràng trên trang.
    """
    path = (SOURCE_DIR / src).resolve()
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.keep_with_next = True
    para.paragraph_format.space_before = Pt(6)
    para.paragraph_format.space_after = Pt(0)
    if not path.is_file():
        add_fragments(para, [Frag(f"[CHÈN HÌNH — {src}]", bold=True, italic=True)])
        return
    from docx.shared import Emu

    run = para.add_run()
    pic = run.add_picture(str(path))
    if pic.width > UIT_TEXT_WIDTH:
        ratio = UIT_TEXT_WIDTH / pic.width
        pic.width, pic.height = Emu(int(pic.width * ratio)), Emu(int(pic.height * ratio))
    # Sơ đồ dọc (lưu đồ tách câu) không được cao quá khoảng hai phần ba trang.
    max_h = Cm(15)
    if pic.height > max_h:
        ratio = max_h / pic.height
        pic.width, pic.height = Emu(int(pic.width * ratio)), Emu(int(pic.height * ratio))


def add_list_of(doc: Document, style_name: str, placeholder: str) -> None:
    field = doc.add_paragraph()
    add_field(field, rf' TOC \h \z \t "{style_name},1" ', placeholder)
    for run in field.runs:
        set_font(run, BODY_FONT)
        run.font.size = Pt(12)
        run.italic = True


def table_font_size(columns: int) -> Pt:
    for limit, size in TABLE_SIZES:
        if columns <= limit:
            return size
    return TABLE_SIZES[-1][1]


def render(
    doc: Document, blocks: list[Block], title_first_heading: bool, toc: bool = False
) -> None:
    if UIT:
        render_uit(doc, blocks)
        return
    first_heading = title_first_heading
    toc_pending = toc
    for block in blocks:
        if block.kind == "directive":
            continue
        if block.kind == "image":
            add_picture_block(doc, (block.lines or [""])[0], block.text)
            continue
        # Mục lục chèn sau phần đầu trang (tiêu đề + dòng thông tin + lời dẫn), tức ngay
        # trước mục đánh số đầu tiên — không phải ở đầu file, trước cả tiêu đề.
        if toc_pending and block.kind == "heading" and block.level >= 2:
            doc.paragraphs[-1].add_run().add_break(WD_BREAK.PAGE)
            add_toc(doc)
            toc_pending = False

        if block.kind == "heading":
            if first_heading and block.level == 1:
                para = doc.add_paragraph()
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                para.paragraph_format.space_after = Pt(14)
                add_fragments(para, inline(block.text, bold=True), size=Pt(18))
                first_heading = False
                continue
            para = doc.add_paragraph(style=f"Heading {min(block.level, 6)}")
            add_fragments(para, inline(block.text, bold=True), HEADING_SIZES[block.level])

        else:
            render_common(doc, block)


def render_common(doc: Document, block: Block) -> None:
    """Khối không phải tiêu đề: giống nhau ở chế độ mặc định và chế độ --uit."""
    if block.kind == "para":
        add_fragments(doc.add_paragraph(), inline(block.text))

    elif block.kind == "math":
        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.space_before = Pt(6)
        para.paragraph_format.space_after = Pt(10)
        add_fragments(para, [Frag(latex_to_text(block.text), italic=True)])

    elif block.kind == "hr":
        # Markdown dùng `---` để ngăn mục; trong Word các Heading đã làm việc đó, thêm
        # đường kẻ chỉ làm trang rối. Bỏ qua có chủ đích.
        return

    elif block.kind == "quote":
        for text in block.lines or []:
            para = doc.add_paragraph()
            pf = para.paragraph_format
            pf.left_indent = Cm(0.8)
            pf.space_after = Pt(6)
            add_fragments(para, inline(text, italic=True))
            add_left_bar(para)

    elif block.kind == "code":
        labels = {"mermaid": "Sơ đồ (mã Mermaid)", "bash": "Lệnh", "json": "JSON"}
        label = labels.get(block.lang)
        if label:
            cap = doc.add_paragraph()
            cap.paragraph_format.space_after = Pt(2)
            add_fragments(cap, [Frag(label, italic=True)], size=Pt(11))
        for text in block.lines or []:
            para = doc.add_paragraph()
            pf = para.paragraph_format
            pf.line_spacing = 1.0
            pf.space_before = pf.space_after = Pt(0)
            pf.left_indent = Cm(0.4)
            pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = para.add_run(text or " ")
            set_font(run, MONO_FONT)
            run.font.size = CODE_SIZE
            shade(para._p.get_or_add_pPr(), SHADE_CODE)
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    elif block.kind == "list":
        for level, marker, text in block.items or []:
            if task := RE_TASK.match(text):
                marker = "☒" if task.group(1).lower() == "x" else "☐"
                text = task.group(2)
            para = doc.add_paragraph()
            pf = para.paragraph_format
            pf.left_indent = Cm(0.7 + 0.6 * level)
            pf.first_line_indent = Cm(-0.5)
            pf.space_after = Pt(3)
            bullet = "•" if marker == "•" and level == 0 else ("–" if marker == "•" else marker)
            add_fragments(para, [Frag(f"{bullet}\t")] + inline(text))

    elif block.kind == "table":
        render_table(doc, block)


RE_REFERENCE = re.compile(r"^\[\d+\]\s")


def render_uit(doc: Document, blocks: list[Block]) -> None:
    """Dựng khóa luận theo quy định của UIT.

    Khác chế độ mặc định ở năm chỗ, đều do quy định đòi:

    * mỗi mục cấp 1 (bìa, lời cảm ơn, danh mục, từng chương) bắt đầu ở trang mới;
    * các trang trước Tóm tắt không đánh số và không vào mục lục — tiêu đề của chúng là
      đoạn chữ canh giữa chứ không phải Heading 1;
    * số trang bắt đầu từ 1 ở Tóm tắt, giữa chân trang;
    * ảnh được chèn thật, chú thích "Hình x.y:" / "Bảng x.y:" mang kiểu riêng để Danh mục
      hình và Danh mục bảng tự dựng được kèm số trang;
    * tài liệu tham khảo thụt treo, canh trái.

    Chỉ thị nằm trong chú thích HTML của file nguồn: ``trang: bìa …`` (sang trang, canh
    giữa như trang bìa), ``trang: không đánh số`` (sang trang), ``trang: bắt đầu đánh số``,
    ``mục lục …``, ``danh-muc: hinh`` và ``danh-muc: bang``.
    """
    cover = False
    numbered = False
    fresh = True  # đang ở đầu một trang mới, chưa có gì
    pending_break = False

    def begin_page() -> None:
        nonlocal pending_break
        if not fresh:
            pending_break = True

    for block in blocks:
        mark = len(doc.paragraphs)

        if block.kind == "directive":
            text = block.text.lower()
            if text.startswith("trang: bìa"):
                begin_page()
                cover = True
            elif text.startswith("trang: không đánh số"):
                begin_page()
                cover = False
            elif text.startswith("trang: bắt đầu đánh số"):
                start_numbering(doc)
                numbered, cover, fresh, pending_break = True, False, True, False
            elif text.startswith("mục lục"):
                begin_page()
                if pending_break:
                    pending_break = False
                    mark = len(doc.paragraphs)
                    add_uit_title(doc, "MỤC LỤC")
                    doc.paragraphs[mark].paragraph_format.page_break_before = True
                else:
                    add_uit_title(doc, "MỤC LỤC")
                field = doc.add_paragraph()
                add_field(field, r' TOC \o "1-3" \h \z \u ', "Mục lục — Word tự dựng khi mở tệp.")
                fresh = False
            elif text.startswith("danh-muc:"):
                kind = "Hình" if "hinh" in text else "Bảng"
                add_list_of(doc, CAPTION_STYLES[kind], f"Danh mục {kind.lower()}.")
                fresh = False
            continue

        if block.kind == "hr":
            continue

        if block.kind == "heading" and block.level == 1:
            begin_page()
            if numbered:
                para = doc.add_paragraph(style="Heading 1")
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_fragments(para, inline(block.text, bold=True), UIT_HEADING_SIZES[1])
            else:
                add_uit_title(doc, block.text)
        elif block.kind == "heading":
            para = doc.add_paragraph(style=f"Heading {min(block.level, 6)}")
            add_fragments(para, inline(block.text, bold=True), UIT_HEADING_SIZES[block.level])
        elif block.kind == "image":
            add_picture_block(doc, (block.lines or [""])[0], block.text)
        elif block.kind == "para" and block.text.strip() == "&nbsp;":
            doc.add_paragraph()
        elif block.kind == "para" and cover:
            para = doc.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.paragraph_format.space_after = Pt(4)
            plain = re.sub(r"[*_]", "", block.text).strip()
            size = next((s for rx, s in COVER_SIZES if rx.search(plain)), Pt(14))
            add_fragments(para, inline(block.text), size=size)
        elif block.kind == "para" and (m := RE_CAPTION.match(block.text)):
            para = doc.add_paragraph(style=CAPTION_STYLES[m.group(1)])
            # Chú thích bảng đứng trên bảng nên phải dính với bảng; chú thích hình đứng
            # dưới hình, còn hình thì đã tự dính xuống chú thích (add_picture_block).
            para.paragraph_format.keep_with_next = m.group(1) == "Bảng"
            add_fragments(para, inline(block.text), size=Pt(12))
        elif block.kind == "para" and RE_REFERENCE.match(block.text):
            para = doc.add_paragraph()
            pf = para.paragraph_format
            pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
            pf.left_indent, pf.first_line_indent = Cm(1), Cm(-1)
            pf.line_spacing = 1.3
            add_fragments(para, inline(block.text), size=Pt(12))
        else:
            render_common(doc, block)

        if pending_break and len(doc.paragraphs) > mark:
            doc.paragraphs[mark].paragraph_format.page_break_before = True
            pending_break = False
        fresh = False


def add_uit_title(doc: Document, text: str) -> None:
    """Tiêu đề trang đầu (lời cảm ơn, danh mục…): canh giữa, đậm, 14 — không vào mục lục."""
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_after = Pt(12)
    para.paragraph_format.keep_with_next = True
    add_fragments(para, inline(text, bold=True), size=Pt(14))


def add_left_bar(para: Paragraph) -> None:
    """Vạch dọc bên trái cho khối trích dẫn."""
    from docx.oxml import OxmlElement

    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "12")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), QUOTE_BAR)
    borders.append(left)
    para._p.get_or_add_pPr().append(borders)


def set_table_borders(table, color: str = "808080", size: str = "6") -> None:
    from docx.oxml import OxmlElement

    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        line = OxmlElement(f"w:{edge}")
        line.set(qn("w:val"), "single")
        line.set(qn("w:sz"), size)
        line.set(qn("w:space"), "0")
        line.set(qn("w:color"), color)
        borders.append(line)
    table._tbl.tblPr.append(borders)


def render_table(doc: Document, block: Block) -> None:
    rows, aligns = block.rows or [], block.aligns or []
    if not rows:
        return
    size = table_font_size(len(rows[0]))
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    # Viền ghi thẳng vào XML chứ không chỉ dựa vào style "Table Grid": trình xem nào
    # không đọc style bảng (Quick Look, một số trình xem trên web) sẽ hiện bảng không kẻ.
    set_table_borders(table)
    align_map = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
    }
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            cell = table.cell(r, c)
            para = cell.paragraphs[0]
            pf = para.paragraph_format
            pf.line_spacing = 1.0
            pf.space_before = pf.space_after = Pt(2)
            pf.alignment = align_map[aligns[c] if c < len(aligns) else "left"]
            add_fragments(para, inline(text, bold=(r == 0)), size=size)
            if r == 0:
                shade(cell._tc.get_or_add_tcPr(), SHADE_TABLE_HEADER)
    if UIT:
        set_column_widths(table, rows)
    # Dòng tiêu đề lặp lại khi bảng tràn sang trang sau.
    from docx.oxml import OxmlElement

    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def set_column_widths(table, rows: list[list[str]]) -> None:
    """Chia bề rộng vùng chữ cho các cột theo lượng chữ của từng cột.

    python-docx tạo bảng với các cột bằng nhau, và Word giữ nguyên như thế: bảng yêu cầu
    chức năng ra cột "Mã" rộng bằng cột nội dung. Chia làm hai bước:

    1. **Sàn** — mỗi cột đủ chứa trọn từ dài nhất của nó. Không có sàn thì cột ngắn bị bóp
       tới mức "PCN01" gãy thành từng ký tự một dòng.
    2. **Phần dư** chia theo lượng chữ trung bình vượt quá sàn, để cột nội dung dài nhận
       phần lớn bề rộng còn lại.
    """
    cols = len(rows[0])
    # Bề rộng trung bình một ký tự Times New Roman ở cỡ chữ của bảng, tính bằng cm — ước
    # rộng tay vì mã và chữ hoa ("PCN01", "WER") rộng hơn chữ thường; chữ đơn cách (trong
    # dấu `) rộng hơn nữa. PAD là lề trong mặc định của ô Word, hai bên cộng lại.
    char_cm = 0.21 * table_font_size(cols).pt / 11
    pad_cm = 0.45

    def cell_cm(text: str) -> tuple[float, float]:
        """(từ dài nhất, cả ô) của một ô, tính bằng cm."""
        mono = set(re.findall(r"`([^`]+)`", text))
        words = re.sub(r"[*_]", "", text).replace("`", "").split()
        # Word ngắt dòng được sau "-" và "/", nên tên mô hình dài như
        # `deepdml/faster-whisper-large-v3-turbo-ct2` chỉ cần chỗ cho đoạn dài nhất.
        pieces = [(piece, w) for w in words for piece in re.split(r"(?<=[-/])", w) if piece]
        def width(piece: str, is_mono: bool) -> float:
            # Chữ hoa và chữ số rộng hơn chữ thường rõ rệt ("WER" gãy thành "WE/R" nếu
            # tính theo bề rộng trung bình); chữ đơn cách thì mọi ký tự rộng như nhau.
            if is_mono:
                return len(piece) * char_cm * 1.2
            return sum(char_cm * (1.3 if ch.isupper() or ch.isdigit() else 1.0) for ch in piece)

        longest = max(
            (width(pc, any(w in m for m in mono)) for pc, w in pieces), default=0
        )
        return longest + pad_cm, len(" ".join(words)) * char_cm + pad_cm

    measured = [[cell_cm(r[c]) for r in rows] for c in range(cols)]
    total = UIT_TEXT_WIDTH.cm
    floor = [max(w for w, _ in col) for col in measured]
    if sum(floor) > total:  # từ quá dài: đành để gãy, nhưng giữ tỷ lệ giữa các cột
        widths_cm = [f * total / sum(floor) for f in floor]
    else:
        want = [
            max(sum(a for _, a in col) / len(col) - f, 0) for col, f in zip(measured, floor)
        ]
        spare = total - sum(floor)
        share = [w / sum(want) if sum(want) else 1 / cols for w in want]
        widths_cm = [f + spare * s for f, s in zip(floor, share)]
    widths = [Cm(w) for w in widths_cm]
    table.autofit = False
    for c, width in enumerate(widths):
        table.columns[c].width = width
        for cell in table.columns[c].cells:
            cell.width = width


def add_toc(doc: Document) -> None:
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_after = Pt(10)
    add_fragments(para, [Frag("MỤC LỤC", bold=True)], size=Pt(14))
    field = doc.add_paragraph()
    add_field(
        field,
        r' TOC \o "1-3" \h \z \u ',
        "Mở file trong Word rồi bấm Ctrl+A, F9 (macOS: Cmd+A, Fn+F9) để dựng mục lục.",
    )
    for run in field.runs:
        set_font(run, BODY_FONT)
        run.font.size = Pt(11)
        run.italic = True
    doc.paragraphs[-1].add_run().add_break(WD_BREAK.PAGE)


def main() -> int:
    ap = argparse.ArgumentParser(description="Markdown trong docs/ → Word (.docx)")
    ap.add_argument("sources", nargs="+", type=Path, help="một hoặc nhiều file .md")
    ap.add_argument("--out", required=True, type=Path, help="file .docx xuất ra")
    ap.add_argument("--toc", action="store_true", help="chèn field mục lục tự động")
    ap.add_argument(
        "--keep-h1",
        action="store_true",
        help="giữ H1 đầu tiên là Heading 1 thay vì dựng thành tiêu đề canh giữa",
    )
    ap.add_argument(
        "--uit",
        action="store_true",
        help="trình bày theo quy định khóa luận của UIT (lề, số trang, ảnh, chú thích)",
    )
    args = ap.parse_args()

    global UIT
    UIT = args.uit

    for src in args.sources:
        if not src.is_file():
            print(f"Không thấy file: {src}", file=sys.stderr)
            return 1

    doc = base_document()

    global SOURCE_DIR
    for n, src in enumerate(args.sources):
        if n:  # mỗi file nguồn bắt đầu ở trang mới
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        SOURCE_DIR = src.parent
        blocks = parse(src.read_text(encoding="utf-8"))
        render(doc, blocks, not args.keep_h1, toc=args.toc and n == 0)

    args.out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.out)
    size_kb = args.out.stat().st_size / 1024
    print(f"{args.out}  ({size_kb:.0f} KB, {len(args.sources)} file nguồn)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
